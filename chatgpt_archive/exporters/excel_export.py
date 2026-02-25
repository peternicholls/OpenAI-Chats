"""Excel exporter for conversation export.

Requires the openpyxl package: pip install chatgpt-archive[excel]
"""

from typing import List

from .base import BaseExporter


class ExcelExporter(BaseExporter):
    """Export conversations to Excel (.xlsx) format.

    Produces an Excel workbook with two sheets:
    - Metadata: Conversation details (title, date, model, etc.)
    - Messages: All messages with role, content, and timestamp columns

    Requires openpyxl to be installed. Install with:
        pip install 'chatgpt-archive[excel]'
    """

    format_name = "excel"
    file_extension = ".xlsx"

    def export(self, conversation: dict, messages: List[dict]) -> str:
        """Export conversation to Excel format.

        Note: Since Excel is a binary format, this returns the workbook
        as base64-encoded content. Use export_bytes() for raw bytes.

        Args:
            conversation: Conversation metadata dictionary
            messages: List of message dictionaries

        Returns:
            Base64-encoded Excel file content

        Raises:
            ImportError: If openpyxl is not installed
        """
        import base64

        raw = self.export_bytes(conversation, messages)
        return base64.b64encode(raw).decode("ascii")

    def export_bytes(self, conversation: dict, messages: List[dict]) -> bytes:
        """Export conversation to Excel format as raw bytes.

        Args:
            conversation: Conversation metadata dictionary
            messages: List of message dictionaries

        Returns:
            Raw Excel file bytes

        Raises:
            ImportError: If openpyxl is not installed
        """
        try:
            from openpyxl import Workbook  # type: ignore
            from openpyxl.styles import Font, Alignment, PatternFill  # type: ignore
        except ImportError:
            raise ImportError(
                "Excel export requires openpyxl. "
                "Install with: pip install 'chatgpt-archive[excel]'"
            )

        import io

        wb = Workbook()

        # --- Metadata Sheet ---
        ws_meta = wb.active
        if ws_meta is None:
            raise RuntimeError("Failed to create worksheet")
        ws_meta.title = "Metadata"

        # Style definitions
        label_font = Font(bold=True)
        header_fill = PatternFill(
            start_color="4472C4", end_color="4472C4", fill_type="solid"
        )
        header_text_font = Font(bold=True, color="FFFFFF", size=11)

        title = self.sanitize_spreadsheet_cell(self.get_display_title(conversation))
        created = self.format_timestamp(conversation.get("create_time"))
        updated = self.format_timestamp(conversation.get("update_time"))
        model = self.sanitize_spreadsheet_cell(conversation.get("model") or "unknown")
        conv_id = self.sanitize_spreadsheet_cell(conversation.get("id", ""))
        msg_count = conversation.get("message_count", len(messages))

        # Write metadata
        metadata_rows = [
            ("Property", "Value"),
            ("Title", title),
            ("Conversation ID", conv_id),
            ("Created", created),
            ("Updated", updated),
            ("Model", model),
            ("Message Count", msg_count),
        ]

        for row_idx, (label, value) in enumerate(metadata_rows, start=1):
            cell_a = ws_meta.cell(row=row_idx, column=1, value=label)
            cell_b = ws_meta.cell(row=row_idx, column=2, value=value)
            if row_idx == 1:
                cell_a.font = header_text_font
                cell_b.font = header_text_font
                cell_a.fill = header_fill
                cell_b.fill = header_fill
            else:
                cell_a.font = label_font

        # Column widths
        ws_meta.column_dimensions["A"].width = 20
        ws_meta.column_dimensions["B"].width = 50

        # --- Messages Sheet ---
        ws_msgs = wb.create_sheet("Messages")

        # Headers
        headers = ["#", "Role", "Content", "Timestamp"]
        for col_idx, header in enumerate(headers, start=1):
            cell = ws_msgs.cell(row=1, column=col_idx, value=header)
            cell.font = header_text_font
            cell.fill = header_fill

        # Role colors
        role_fills = {
            "user": PatternFill(
                start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
            ),
            "assistant": PatternFill(
                start_color="DAEEF3", end_color="DAEEF3", fill_type="solid"
            ),
            "system": PatternFill(
                start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"
            ),
            "tool": PatternFill(
                start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"
            ),
        }

        # Write messages
        row_num = 1
        for msg in messages:
            role = msg.get("role", "unknown")
            content = msg.get("content") or ""

            # Skip empty system messages
            if not content.strip() and role == "system":
                continue

            row_num += 1
            timestamp = self.format_timestamp(msg.get("create_time"))
            fill = role_fills.get(role)
            safe_role = self.sanitize_spreadsheet_cell(role.capitalize())
            safe_content = self.sanitize_spreadsheet_cell(content)
            safe_timestamp = self.sanitize_spreadsheet_cell(timestamp)

            cells = [
                ws_msgs.cell(row=row_num, column=1, value=row_num - 1),
                ws_msgs.cell(row=row_num, column=2, value=safe_role),
                ws_msgs.cell(row=row_num, column=3, value=safe_content),
                ws_msgs.cell(row=row_num, column=4, value=safe_timestamp),
            ]

            # Apply role-based fill color
            if fill:
                for cell in cells:
                    cell.fill = fill

            # Wrap text for content column
            cells[2].alignment = Alignment(wrap_text=True, vertical="top")

        # Column widths
        ws_msgs.column_dimensions["A"].width = 6
        ws_msgs.column_dimensions["B"].width = 12
        ws_msgs.column_dimensions["C"].width = 80
        ws_msgs.column_dimensions["D"].width = 22

        # Freeze header row
        ws_msgs.freeze_panes = "A2"

        # Save to bytes
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
